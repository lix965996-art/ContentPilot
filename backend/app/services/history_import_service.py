"""CSV / XLSX import of historical engagement records.

The importer is deliberately mapping-driven: platform exports and public
research datasets use different column names, so the caller previews a
suggested mapping, adjusts it, and only then commits.
"""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.exceptions import AppException
from app.models.business import EngagementHistory, HistoryImportBatch, PlatformAccount
from app.services.activity_service import (
    SOURCE_TYPE_LABELS,
    compute_metrics,
    content_type_name,
    normalize_content_type,
    parse_publish_time,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PREVIEW_ROWS = 20

CANONICAL_FIELDS: list[dict[str, Any]] = [
    {"field": "platform", "label": "平台", "required": True, "type": "text"},
    {"field": "account_id", "label": "账号标识", "required": False, "type": "text"},
    {"field": "content_type", "label": "内容类型", "required": False, "type": "text"},
    {"field": "publish_time", "label": "发布时间", "required": True, "type": "datetime"},
    {"field": "views", "label": "浏览量", "required": False, "type": "number"},
    {"field": "impressions", "label": "曝光量", "required": False, "type": "number"},
    {"field": "likes", "label": "点赞", "required": False, "type": "number"},
    {"field": "comments", "label": "评论", "required": False, "type": "number"},
    {"field": "shares", "label": "转发/分享", "required": False, "type": "number"},
    {"field": "favorites", "label": "收藏", "required": False, "type": "number"},
    {"field": "followers", "label": "粉丝数", "required": False, "type": "number"},
    {"field": "source_type", "label": "数据来源", "required": False, "type": "text"},
]
REQUIRED_FIELDS = [item["field"] for item in CANONICAL_FIELDS if item["required"]]
NUMERIC_FIELDS = [item["field"] for item in CANONICAL_FIELDS if item["type"] == "number"]
# At least one interaction column must be present, otherwise the row cannot
# contribute to any engagement statistic.
INTERACTION_FIELDS = ["likes", "comments", "shares", "favorites"]

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "platform": ("platform", "平台", "channel", "site", "source_platform"),
    "account_id": (
        "account_id",
        "accountid",
        "account",
        "账号",
        "账号id",
        "channel_id",
        "channelid",
        "author",
        "uid",
    ),
    "content_type": (
        "content_type",
        "contenttype",
        "内容类型",
        "category",
        "categoryid",
        "type",
        "topic",
    ),
    "publish_time": (
        "publish_time",
        "publishtime",
        "发布时间",
        "published_at",
        "publishedat",
        "post_time",
        "created_at",
        "date",
        "datetime",
        "upload_date",
    ),
    "views": ("views", "浏览量", "view_count", "viewcount", "play_count", "reads", "阅读量"),
    "impressions": ("impressions", "曝光量", "impression", "reach", "exposure", "展现量"),
    "likes": ("likes", "点赞", "like_count", "likecount", "digg_count", "赞"),
    "comments": ("comments", "评论", "comment_count", "commentcount", "replies", "评论数"),
    "shares": ("shares", "转发", "分享", "share_count", "sharecount", "reposts", "forwards"),
    "favorites": (
        "favorites",
        "收藏",
        "favorite_count",
        "favoritecount",
        "collects",
        "collect_count",
        "bookmarks",
        "saves",
    ),
    "followers": ("followers", "粉丝", "粉丝数", "follower_count", "subscribers", "fans"),
    "source_type": ("source_type", "sourcetype", "数据来源", "source", "dataset"),
}

SOURCE_TYPE_CHOICES = list(SOURCE_TYPE_LABELS.keys())


def _normalize_header(value: str) -> str:
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def read_table(raw: bytes, filename: str) -> tuple[list[str], list[dict[str, Any]]]:
    """Parse an uploaded CSV/XLSX file into headers plus raw string records."""
    name = (filename or "").lower()
    if not name.endswith((".csv", ".xlsx")):
        raise AppException(40041, "只支持 CSV 或 XLSX 文件")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise AppException(40042, "文件不能超过 10MB")
    if name.endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gbk")
            except UnicodeDecodeError as exc:
                raise AppException(40043, "CSV 需要使用 UTF-8 或 GBK 编码") from exc
        reader = csv.DictReader(io.StringIO(text, newline=""))
        headers = [str(item) for item in (reader.fieldnames or [])]
        records = [dict(row) for row in reader]
    else:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        headers = [str(item) if item is not None else "" for item in (rows[0] if rows else [])]
        records = [dict(zip(headers, values, strict=False)) for values in rows[1:]]
        workbook.close()
    if not headers:
        raise AppException(40045, "文件没有表头，无法识别字段")
    return headers, records


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Guess which uploaded column feeds each canonical field."""
    normalized = {_normalize_header(header): header for header in headers if header}
    mapping: dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            key = _normalize_header(alias)
            if key in normalized:
                mapping[field] = normalized[key]
                break
        else:
            mapping[field] = ""
    return mapping


def _to_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    text = str(value).strip().replace(",", "").replace("，", "")
    if not text:
        return 0
    lowered = text.lower()
    multiplier = 1
    if lowered.endswith("w") or lowered.endswith("万"):
        multiplier, text = 10_000, text[:-1]
    elif lowered.endswith("k"):
        multiplier, text = 1_000, text[:-1]
    try:
        return max(0, int(round(float(text) * multiplier)))
    except ValueError as exc:
        raise ValueError(f"无法解析数值：{value}") from exc


def row_hash(
    platform: str, account_ref: str, publish_time: datetime, metrics: dict[str, int]
) -> str:
    payload = "|".join(
        [
            platform,
            account_ref,
            publish_time.isoformat(timespec="minutes"),
            *[f"{key}={metrics.get(key, 0)}" for key in sorted(NUMERIC_FIELDS)],
        ]
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def parse_rows(
    records: list[dict[str, Any]],
    mapping: dict[str, str],
    *,
    default_platform: str = "",
    default_source_type: str = "ACCOUNT_HISTORY",
    default_content_type: str = "",
) -> dict[str, Any]:
    """Validate and normalise every record; never raises on a single bad row."""
    missing_required = [
        field
        for field in REQUIRED_FIELDS
        if not mapping.get(field) and not (field == "platform" and default_platform)
    ]
    if missing_required:
        raise AppException(
            40044,
            "必填字段未完成映射",
            422,
            {"missingFields": missing_required},
        )

    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[str] = set()
    duplicate_rows = 0
    missing_value_rows = 0

    for index, record in enumerate(records, start=2):
        try:
            platform = str(record.get(mapping.get("platform", ""), "") or default_platform).strip()
            platform = platform.upper().replace(" ", "_") if platform else ""
            if not platform:
                raise ValueError("platform 不能为空")
            publish_time = parse_publish_time(record.get(mapping.get("publish_time", "")))
            metrics = {
                field: _to_int(record.get(mapping.get(field, ""))) for field in NUMERIC_FIELDS
            }
            if not any(metrics[field] for field in INTERACTION_FIELDS):
                raise ValueError("点赞、评论、分享、收藏至少需要一项有效数据")
            account_ref = str(record.get(mapping.get("account_id", ""), "") or "").strip()
            content_type = normalize_content_type(
                str(record.get(mapping.get("content_type", ""), "") or default_content_type)
            )
            source_type = (
                str(record.get(mapping.get("source_type", ""), "") or default_source_type)
                .strip()
                .upper()
                .replace(" ", "_")
            )
            if source_type not in SOURCE_TYPE_CHOICES:
                source_type = default_source_type
            computed = compute_metrics(
                views=metrics["views"],
                impressions=metrics["impressions"],
                likes=metrics["likes"],
                comments=metrics["comments"],
                shares=metrics["shares"],
                favorites=metrics["favorites"],
                followers=metrics["followers"],
            )
            if computed["reach"] == 0:
                missing_value_rows += 1
            digest = row_hash(platform, account_ref, publish_time, metrics)
            duplicate_in_file = digest in seen
            if duplicate_in_file:
                duplicate_rows += 1
            seen.add(digest)
            parsed.append(
                {
                    "row": index,
                    "platform": platform,
                    "accountRef": account_ref,
                    "contentType": content_type,
                    "publishTime": publish_time,
                    "dayOfWeek": publish_time.weekday(),
                    "hourOfDay": publish_time.hour,
                    "sourceType": source_type,
                    "rowHash": digest,
                    "duplicateInFile": duplicate_in_file,
                    **metrics,
                    **computed,
                }
            )
        except (ValueError, TypeError, KeyError) as exc:
            errors.append({"row": index, "message": str(exc)[:200]})

    return {
        "rows": parsed,
        "errors": errors,
        "totalRows": len(records),
        "validRows": len(parsed),
        "errorRows": len(errors),
        "duplicateRows": duplicate_rows,
        "missingValueRows": missing_value_rows,
    }


def _serialize_preview(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "row": item["row"],
        "platform": item["platform"],
        "accountRef": item["accountRef"],
        "contentType": item["contentType"],
        "contentTypeName": content_type_name(item["contentType"]),
        "status": "文件内重复" if item["duplicateInFile"] else "可导入",
        "publishTime": item["publishTime"].isoformat(),
        "dayOfWeek": item["dayOfWeek"],
        "hourOfDay": item["hourOfDay"],
        "views": item["views"],
        "impressions": item["impressions"],
        "likes": item["likes"],
        "comments": item["comments"],
        "shares": item["shares"],
        "favorites": item["favorites"],
        "followers": item["followers"],
        "sourceType": item["sourceType"],
        "engagementRate": item["engagementRate"],
        "metricBasis": item["metricBasis"],
        "duplicateInFile": item["duplicateInFile"],
    }


def build_preview(
    db: Session,
    headers: list[str],
    records: list[dict[str, Any]],
    mapping: dict[str, str],
    *,
    default_platform: str = "",
    default_source_type: str = "ACCOUNT_HISTORY",
    default_content_type: str = "",
) -> dict[str, Any]:
    result = parse_rows(
        records,
        mapping,
        default_platform=default_platform,
        default_source_type=default_source_type,
        default_content_type=default_content_type,
    )
    hashes = [item["rowHash"] for item in result["rows"]]
    existing: set[str] = set()
    for start in range(0, len(hashes), 500):
        chunk = hashes[start : start + 500]
        existing.update(
            db.scalars(
                select(EngagementHistory.row_hash).where(EngagementHistory.row_hash.in_(chunk))
            ).all()
        )
    duplicate_in_db = sum(1 for item in result["rows"] if item["rowHash"] in existing)
    importable = sum(
        1
        for item in result["rows"]
        if not item["duplicateInFile"] and item["rowHash"] not in existing
    )
    return {
        "headers": headers,
        "mapping": mapping,
        "fields": CANONICAL_FIELDS,
        "totalRows": result["totalRows"],
        "validRows": result["validRows"],
        "errorRows": result["errorRows"],
        "duplicateInFile": result["duplicateRows"],
        "duplicateInDatabase": duplicate_in_db,
        "missingValueRows": result["missingValueRows"],
        "importableRows": importable,
        "errors": result["errors"][:50],
        "preview": [_serialize_preview(item) for item in result["rows"][:MAX_PREVIEW_ROWS]],
        "sourceTypes": [
            {"value": key, "label": value} for key, value in SOURCE_TYPE_LABELS.items()
        ],
    }


def commit_import(
    db: Session,
    *,
    filename: str,
    headers: list[str],
    records: list[dict[str, Any]],
    mapping: dict[str, str],
    user_id: int | None,
    default_platform: str = "",
    default_source_type: str = "ACCOUNT_HISTORY",
    default_content_type: str = "",
    source_note: str = "",
) -> HistoryImportBatch:
    result = parse_rows(
        records,
        mapping,
        default_platform=default_platform,
        default_source_type=default_source_type,
        default_content_type=default_content_type,
    )
    accounts = {
        account.platform: account.id for account in db.scalars(select(PlatformAccount)).all()
    }
    batch = HistoryImportBatch(
        filename=filename[:255],
        source_type=default_source_type,
        source_note=(source_note or SOURCE_TYPE_LABELS.get(default_source_type, ""))[:255],
        platform_hint=default_platform or None,
        total_rows=result["totalRows"],
        field_mapping_json={"mapping": mapping, "headers": headers},
        status="COMPLETED",
        created_by=user_id,
    )
    db.add(batch)
    db.flush()

    existing = set(
        db.scalars(
            select(EngagementHistory.row_hash).where(
                EngagementHistory.row_hash.in_([item["rowHash"] for item in result["rows"]])
            )
        ).all()
    )
    inserted = 0
    duplicates = 0
    for item in result["rows"]:
        if item["duplicateInFile"] or item["rowHash"] in existing:
            duplicates += 1
            continue
        existing.add(item["rowHash"])
        db.add(
            EngagementHistory(
                batch_id=batch.id,
                platform=item["platform"],
                account_ref=item["accountRef"][:120],
                account_id=accounts.get(item["platform"]),
                content_type=item["contentType"],
                publish_time=item["publishTime"],
                day_of_week=item["dayOfWeek"],
                hour_of_day=item["hourOfDay"],
                views=item["views"],
                impressions=item["impressions"],
                likes=item["likes"],
                comments=item["comments"],
                shares=item["shares"],
                favorites=item["favorites"],
                followers=item["followers"],
                engagement_total=item["engagementTotal"],
                engagement_rate=item["engagementRate"],
                # ``engagement_score`` stores the saturated interaction-rate
                # component; the view component is normalised per corpus at
                # analysis time (see activity_service.composite_score).
                engagement_score=item["rateComponent"],
                metric_detail_json={
                    "likeRate": item["likeRate"],
                    "commentRate": item["commentRate"],
                    "shareRate": item["shareRate"],
                    "favoriteRate": item["favoriteRate"],
                    "metricBasis": item["metricBasis"],
                    "reach": item["reach"],
                },
                source_type=item["sourceType"],
                source_note=(source_note or SOURCE_TYPE_LABELS.get(item["sourceType"], ""))[:255],
                row_hash=item["rowHash"],
                created_by=user_id,
            )
        )
        inserted += 1

    batch.success_count = inserted
    batch.duplicate_count = duplicates
    batch.error_count = result["errorRows"]
    batch.missing_value_count = result["missingValueRows"]
    batch.errors_json = result["errors"][:100]
    if inserted == 0:
        batch.status = "NO_NEW_DATA" if result["rows"] else "FAILED"
    db.flush()
    return batch


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "历史互动数据"
    sheet.append([item["field"] for item in CANONICAL_FIELDS])
    sheet.append(
        [
            "WEIBO",
            "demo-account",
            "KNOWLEDGE",
            "2026-07-01 20:00",
            12000,
            15000,
            320,
            48,
            26,
            64,
            8600,
            "ACCOUNT_HISTORY",
        ]
    )
    notes = workbook.create_sheet("字段说明")
    notes.append(["字段", "含义", "是否必填"])
    for item in CANONICAL_FIELDS:
        notes.append([item["field"], item["label"], "必填" if item["required"] else "可选"])
    notes.append(["source_type", "可选值：" + " / ".join(SOURCE_TYPE_CHOICES), "可选"])
    notes.append(
        [
            "说明",
            "YouTube 公开样本请填写 YOUTUBE_PUBLIC_SAMPLE，"
            "系统会标注为公开样本，不会冒充国内平台数据。",
            "",
        ]
    )
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
